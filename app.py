from flask import Flask, render_template, request, jsonify
import io
import os
import re
import random
from collections import Counter

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    from PyPDF2 import PdfReader
except ImportError:
    PdfReader = None

try:
    from docx import Document
except ImportError:
    Document = None


app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024


# ============================================================
# SAMPLE LECTURE
# ============================================================

SAMPLE_LECTURE = """
Artificial Intelligence is a branch of computer science that focuses on
creating systems capable of performing tasks that normally require human
intelligence. These tasks include learning, reasoning, problem solving,
understanding natural language, recognizing patterns, and making decisions.

Machine Learning is a major part of Artificial Intelligence. It allows
computers to learn from data without being explicitly programmed for every
task. Machine learning algorithms identify patterns in data and use those
patterns to make predictions or decisions.

Deep Learning is a specialized area of machine learning that uses artificial
neural networks with multiple layers. Deep learning is widely used in image
recognition, speech recognition, natural language processing, and autonomous
systems.

Natural Language Processing allows computers to understand, process, and
generate human language. It is used in chatbots, translation systems,
sentiment analysis, search engines, and voice assistants.

Computer Vision enables computers to understand information from images and
videos. It can be used for object detection, facial recognition, medical
image analysis, and autonomous vehicles.

Artificial Intelligence has many applications in education, healthcare,
finance, transportation, manufacturing, and entertainment. AI systems can
automate repetitive tasks, support decision making, and analyze large amounts
of information quickly.

However, AI also creates challenges related to privacy, security, bias,
employment, and ethical decision making. Responsible development of AI
requires careful consideration of these issues.
"""


# ============================================================
# TEXT CLEANING
# ============================================================

def normalize_uploaded_text(text):
    """
    Clean uploaded text while preserving its logical order.
    """

    if not text:
        return ""

    text = text.replace("\u00ad", "")
    text = text.replace("\r\n", "\n")
    text = text.replace("\r", "\n")
    text = text.replace("\t", " ")

    # Fix words broken across PDF lines:
    # informa-
    # tion
    # becomes:
    # information
    text = re.sub(
        r"(?<=[A-Za-z])-\n(?=[A-Za-z])",
        "",
        text
    )

    raw_lines = text.split("\n")
    lines = []

    for line in raw_lines:
        line = re.sub(r"[ \t]+", " ", line).strip()
        lines.append(line)

    paragraphs = []
    current = []

    for line in lines:

        if line:
            current.append(line)

        else:
            if current:
                paragraphs.append(" ".join(current))
                current = []

    if current:
        paragraphs.append(" ".join(current))

    result = "\n\n".join(paragraphs)

    result = re.sub(
        r"\n{3,}",
        "\n\n",
        result
    )

    return result.strip()


# ============================================================
# FILE EXTRACTION
# ============================================================

def extract_txt(data):

    for encoding in (
        "utf-8-sig",
        "utf-8",
        "utf-16",
        "latin-1"
    ):
        try:
            return normalize_uploaded_text(
                data.decode(encoding)
            )
        except UnicodeDecodeError:
            continue

    return normalize_uploaded_text(
        data.decode(
            "utf-8",
            errors="ignore"
        )
    )


def extract_docx(data):

    if Document is None:
        raise RuntimeError(
            "DOCX support is missing. Run: pip install python-docx"
        )

    document = Document(
        io.BytesIO(data)
    )

    paragraphs = []

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    return normalize_uploaded_text(
        "\n\n".join(paragraphs)
    )


def extract_pdf(data):

    pages = []

    # --------------------------------------------------------
    # PyMuPDF
    # --------------------------------------------------------

    if fitz is not None:

        try:

            document = fitz.open(
                stream=data,
                filetype="pdf"
            )

            for page in document:

                page_text = page.get_text(
                    "text",
                    sort=True
                )

                if page_text.strip():
                    pages.append(page_text)

            document.close()

            if pages:
                return normalize_uploaded_text(
                    "\n\n".join(pages)
                )

        except Exception:
            pass

    # --------------------------------------------------------
    # PyPDF2 fallback
    # --------------------------------------------------------

    if PdfReader is not None:

        try:

            reader = PdfReader(
                io.BytesIO(data)
            )

            for page in reader.pages:

                try:
                    page_text = page.extract_text() or ""
                except Exception:
                    page_text = ""

                if page_text.strip():
                    pages.append(page_text)

            if pages:
                return normalize_uploaded_text(
                    "\n\n".join(pages)
                )

        except Exception:
            pass

    raise RuntimeError(
        "Unable to read this PDF. "
        "Make sure the PDF contains selectable text."
    )


def extract_uploaded_file(file):

    filename = file.filename or ""

    extension = os.path.splitext(
        filename
    )[1].lower()

    data = file.read()

    if not data:
        raise ValueError(
            "The uploaded file is empty."
        )

    if extension in [".txt", ".md"]:
        return extract_txt(data)

    if extension == ".pdf":
        return extract_pdf(data)

    if extension == ".docx":
        return extract_docx(data)

    raise ValueError(
        "Unsupported file type. "
        "Use PDF, DOCX, TXT, or MD."
    )


# ============================================================
# SENTENCE SPLITTING
# ============================================================

def split_sentences(text):

    if not text:
        return []

    text = re.sub(
        r"\s+",
        " ",
        text
    ).strip()

    if not text:
        return []

    # Protect common abbreviations.
    abbreviations = [
        "e.g.",
        "i.e.",
        "etc.",
        "Mr.",
        "Mrs.",
        "Dr.",
        "Prof.",
        "Fig.",
        "No.",
        "vs."
    ]

    placeholders = {}

    for i, abbreviation in enumerate(abbreviations):

        placeholder = f"__ABBR_{i}__"

        text = text.replace(
            abbreviation,
            placeholder
        )

        placeholders[placeholder] = abbreviation

    parts = re.split(
        r"(?<=[.!?])\s+(?=[A-Z0-9])",
        text
    )

    sentences = []

    for part in parts:

        for placeholder, abbreviation in placeholders.items():
            part = part.replace(
                placeholder,
                abbreviation
            )

        part = part.strip()

        if part:
            sentences.append(part)

    return sentences


# ============================================================
# IMPORTANT LONG FILE FIX
# ============================================================

def build_content_sections(text):

    """
    Divides long documents into logical ordered sections.

    This prevents a long uploaded file from becoming one
    huge mixed paragraph.
    """

    if not text:
        return []

    paragraphs = [
        paragraph.strip()
        for paragraph in re.split(
            r"\n\s*\n",
            text
        )
        if paragraph.strip()
    ]

    sections = []

    for paragraph in paragraphs:

        sentences = split_sentences(
            paragraph
        )

        if not sentences:
            continue

        # Keep short paragraphs together.
        if len(sentences) <= 5:

            sections.append(
                " ".join(sentences)
            )

            continue

        # Split long paragraphs into groups.
        current = []

        for sentence in sentences:

            current.append(sentence)

            if len(current) >= 4:

                sections.append(
                    " ".join(current)
                )

                current = []

        if current:

            sections.append(
                " ".join(current)
            )

    return sections


# ============================================================
# KEYWORDS
# ============================================================

STOP_WORDS = {
    "the",
    "and",
    "that",
    "this",
    "with",
    "from",
    "have",
    "has",
    "are",
    "was",
    "were",
    "will",
    "would",
    "could",
    "should",
    "their",
    "there",
    "they",
    "them",
    "then",
    "than",
    "into",
    "about",
    "also",
    "which",
    "when",
    "where",
    "what",
    "while",
    "using",
    "used",
    "use",
    "such",
    "these",
    "those",
    "through",
    "between",
    "being",
    "been",
    "because",
    "more",
    "most",
    "some",
    "many",
    "each",
    "other",
    "only",
    "very",
    "often",
    "can",
    "may",
    "might",
    "must",
    "not",
    "for",
    "you",
    "your",
    "our",
    "its",
    "his",
    "her",
    "but",
    "how",
    "why",
    "who",
    "whose",
    "all",
    "any",
    "both",
    "one",
    "two",
    "first",
    "second",
    "new",
    "different",
    "important",
    "system",
    "systems"
}


def extract_keywords(text, limit=12):

    words = re.findall(
        r"\b[A-Za-z][A-Za-z0-9'-]{3,}\b",
        text.lower()
    )

    words = [
        word
        for word in words
        if word not in STOP_WORDS
    ]

    frequency = Counter(words)

    return [
        word
        for word, count in frequency.most_common()
    ][:limit]


# ============================================================
# SENTENCE SCORE
# ============================================================

def sentence_score(
    sentence,
    keywords
):

    words = set(
        re.findall(
            r"\b[A-Za-z][A-Za-z0-9'-]{3,}\b",
            sentence.lower()
        )
    )

    keyword_score = sum(
        1
        for word in words
        if word in keywords
    )

    length_score = min(
        len(words) / 25,
        1.0
    )

    return (
        keyword_score +
        length_score
    )


# ============================================================
# SUMMARY
# ============================================================

def make_summary(
    text,
    sections
):

    if not text:
        return ""

    if not sections:
        sections = [text]

    # Short text.
    if len(sections) <= 2:

        sentences = split_sentences(
            text
        )

        return " ".join(
            sentences[:5]
        )

    keywords = set(
        extract_keywords(
            text,
            limit=30
        )
    )

    candidates = []

    for index, section in enumerate(sections):

        sentences = split_sentences(
            section
        )

        if not sentences:
            continue

        best_sentence = max(
            sentences,
            key=lambda sentence:
                sentence_score(
                    sentence,
                    keywords
                )
        )

        candidates.append({
            "index": index,
            "sentence": best_sentence,
            "score": sentence_score(
                best_sentence,
                keywords
            )
        })

    if not candidates:
        return text[:1200]

    # Select points from across the document.
    target_count = min(
        8,
        max(
            3,
            len(candidates) // 3
        )
    )

    if len(candidates) <= target_count:

        selected = candidates

    else:

        selected = []

        step = (
            len(candidates) /
            target_count
        )

        for i in range(target_count):

            start = int(
                i * step
            )

            end = int(
                (i + 1) * step
            )

            region = candidates[
                start:end
            ]

            if region:

                selected.append(
                    max(
                        region,
                        key=lambda item:
                            item["score"]
                    )
                )

    # VERY IMPORTANT:
    # Put selected summary sentences back
    # into original document order.
    selected.sort(
        key=lambda item:
            item["index"]
    )

    return " ".join(
        item["sentence"]
        for item in selected
    )


# ============================================================
# KEY POINTS
# ============================================================

def make_key_points(
    text,
    sections,
    maximum=10
):

    if not text:
        return []

    if not sections:
        sections = [text]

    keywords = set(
        extract_keywords(
            text,
            limit=40
        )
    )

    candidates = []

    for index, section in enumerate(sections):

        sentences = split_sentences(
            section
        )

        scored = []

        for sentence in sentences:

            if len(
                sentence.split()
            ) < 6:
                continue

            score = sentence_score(
                sentence,
                keywords
            )

            scored.append(
                (
                    score,
                    sentence
                )
            )

        if scored:

            score, sentence = max(
                scored,
                key=lambda item:
                    item[0]
            )

            candidates.append({
                "index": index,
                "score": score,
                "text": sentence
            })

    if not candidates:
        return []

    target = min(
        maximum,
        len(candidates)
    )

    if len(candidates) <= target:

        selected = candidates

    else:

        selected = []

        step = (
            len(candidates) /
            target
        )

        for i in range(target):

            start = int(
                i * step
            )

            end = int(
                (i + 1) * step
            )

            region = candidates[
                start:end
            ]

            if region:

                selected.append(
                    max(
                        region,
                        key=lambda item:
                            item["score"]
                    )
                )

    # Restore original order.
    selected.sort(
        key=lambda item:
            item["index"]
    )

    points = []

    for item in selected:

        point = item["text"].strip()

        point = re.sub(
            r"^[•●▪◦\-*]+\s*",
            "",
            point
        )

        if (
            point
            and point not in points
        ):
            points.append(point)

    return points[:maximum]


# ============================================================
# NOTES
# ============================================================

def make_notes(
    text,
    sections
):

    summary = make_summary(
        text,
        sections
    )

    points = make_key_points(
        text,
        sections
    )

    keywords = extract_keywords(
        text,
        limit=12
    )

    return {
        "overview": summary,
        "key_points": points,
        "keywords": keywords
    }


# ============================================================
# QUIZ
# ============================================================

def make_quiz(
    text,
    sections,
    count=8
):

    if not sections:
        sections = [text]

    questions = []

    # Distribute questions through the document.
    if len(sections) > count:

        selected_sections = []

        step = (
            len(sections) /
            count
        )

        for i in range(count):

            index = min(
                int(i * step),
                len(sections) - 1
            )

            selected_sections.append(
                sections[index]
            )

    else:

        selected_sections = sections

    all_keywords = extract_keywords(
        text,
        limit=40
    )

    for section in selected_sections:

        sentences = split_sentences(
            section
        )

        if not sentences:
            continue

        sentence = max(
            sentences,
            key=lambda item:
                len(item.split())
        )

        local_keywords = extract_keywords(
            sentence,
            limit=8
        )

        if not local_keywords:
            continue

        answer = local_keywords[0]

        distractors = [
            word
            for word in all_keywords
            if word != answer
        ]

        distractors = list(
            dict.fromkeys(
                distractors
            )
        )

        if len(distractors) < 3:
            continue

        options = [
            answer,
            distractors[0],
            distractors[1],
            distractors[2]
        ]

        random.shuffle(options)

        correct_index = options.index(
            answer
        )

        questions.append({
            "question":
                f'Which term is most closely related to this statement?\n\n"{sentence}"',

            "options": options,

            "answer": correct_index,

            "explanation": sentence
        })

        if len(questions) >= count:
            break

    return questions


# ============================================================
# FLASHCARDS
# ============================================================

def make_flashcards(
    text,
    sections,
    count=8
):

    cards = []

    for section in sections:

        keywords = extract_keywords(
            section,
            limit=8
        )

        if not keywords:
            continue

        keyword = keywords[0]

        cards.append({
            "question":
                f"Explain: {keyword}",

            "answer":
                section
        })

        if len(cards) >= count:
            break

    return cards


# ============================================================
# TITLE
# ============================================================

def make_title(text):

    keywords = extract_keywords(
        text,
        limit=3
    )

    if keywords:

        return " • ".join(
            word.title()
            for word in keywords
        )

    return "Lecture Notes"


# ============================================================
# MAIN GENERATION
# ============================================================

def generate_content(text):

    text = normalize_uploaded_text(
        text
    )

    if not text:
        raise ValueError(
            "Please enter lecture text or upload a file."
        )

    # --------------------------------------------------------
    # LONG FILE FIX
    # --------------------------------------------------------

    sections = build_content_sections(
        text
    )

    if not sections:
        sections = [text]

    # --------------------------------------------------------
    # Generate everything from the same
    # properly ordered sections.
    # --------------------------------------------------------

    notes = make_notes(
        text,
        sections
    )

    quiz = make_quiz(
        text,
        sections
    )

    flashcards = make_flashcards(
        text,
        sections
    )

    return {
        "title": make_title(text),

        "summary": notes["overview"],

        "notes": notes,

        "quiz": quiz,

        "flashcards": flashcards,

        "source_text": text
    }


# ============================================================
# HOME
# ============================================================

@app.route("/")
def index():

    return render_template(
        "index.html"
    )


# ============================================================
# GENERATE
# ============================================================

@app.route(
    "/generate",
    methods=["POST"]
)
def generate():

    try:

        data = request.get_json(
            silent=True
        ) or {}

        text = data.get(
            "text",
            ""
        )

        result = generate_content(
            text
        )

        # IMPORTANT:
        # Return fields DIRECTLY.
        #
        # Your existing JavaScript expects:
        # response.summary
        # response.notes
        # response.quiz
        # response.flashcards
        #
        # NOT:
        # response.data.summary
        #
        # We also keep "data" for compatibility.
        # ----------------------------------------------------

        response = dict(result)

        response["success"] = True
        response["data"] = result

        return jsonify(response)

    except Exception as error:

        print(
            "GENERATION ERROR:",
            error
        )

        return jsonify({
            "success": False,
            "error": str(error)
        }), 400


# ============================================================
# UPLOAD
# ============================================================

@app.route(
    "/upload",
    methods=["POST"]
)
def upload():

    try:

        if "file" not in request.files:

            return jsonify({
                "success": False,
                "error":
                    "No file was uploaded."
            }), 400

        file = request.files["file"]

        if not file.filename:

            return jsonify({
                "success": False,
                "error":
                    "Please select a file."
            }), 400

        text = extract_uploaded_file(
            file
        )

        if not text.strip():

            return jsonify({
                "success": False,
                "error":
                    "No readable text was found in the file."
            }), 400

        # Return text directly because the
        # existing frontend expects response.text.
        return jsonify({
            "success": True,
            "text": text,
            "filename": file.filename
        })

    except Exception as error:

        print(
            "UPLOAD ERROR:",
            error
        )

        return jsonify({
            "success": False,
            "error": str(error)
        }), 400


# ============================================================
# HEALTH CHECK
# ============================================================

@app.route("/health")
def health():

    return jsonify({
        "status": "ok"
    })


# ============================================================
# ERROR HANDLING
# ============================================================

@app.errorhandler(413)
def file_too_large(error):

    return jsonify({
        "success": False,
        "error":
            "File is too large. Maximum size is 10 MB."
    }), 413


@app.errorhandler(500)
def server_error(error):

    return jsonify({
        "success": False,
        "error":
            "Something went wrong on the server."
    }), 500


# ============================================================
# START SERVER
# ============================================================

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5001)